"""Compare earlier QUASAR ARDs with W188 without merging. pip install openpyxl"""
from pathlib import Path
from collections import Counter
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
import csv, json, sqlite3, tempfile, zlib, hashlib, os
from openpyxl import load_workbook

INPUT_FOLDER = Path(r'C:\Users\JMende95\OneDrive - JNJ\Desktop\ard_data')
REFERENCE = 'CNTO1959UCO3001_quasar_lte_wk188_ard_20261013.xlsx'
EARLIER = [
 'CNTO1959UCO3001_quasar_ind1_wk32_ard_20260619.xlsx',
 'CNTO1959UCO3001_quasar_ind2_wk32_ard_20260621.xlsx',
 'CNTO1959UCO3001_quasar_maint_wk44_ard_20260616.xlsx',
 'CNTO1959UCO3001_quasar_lte_wk92_ard_20261012.xlsx',
 'CNTO1959UCO3001_quasar_lte_wk140_ard_20261013.xlsx',
]
KEYS = ('USUBJID','AVISIT','AVISITN','AVISIT_ORDER')
OUTPUT_FOLDER = INPUT_FOLDER / 'retention_audit'
BATCH = 250

def canonical(v):
    if v is None: return None
    if isinstance(v, str):
        s=v.strip()
        if not s: return None
        return s
    if isinstance(v, (datetime,date)): return v.isoformat(sep=' ') if isinstance(v,datetime) else v.isoformat()
    if isinstance(v,bool): return str(v)
    try:
        n=Decimal(str(v))
        if n.is_finite(): return format(n.normalize(),'f')
    except (InvalidOperation,ValueError): pass
    return str(v).strip()

def key_for(row, indices):
    vals=[canonical(row[i]) if i<len(row) else None for i in indices]
    return json.dumps(vals,ensure_ascii=False,separators=(',',':')) if all(x is not None for x in vals) else None

def rows_from(path):
    wb=load_workbook(path,read_only=True,data_only=True)
    try:
        if 'ARD' not in wb.sheetnames: raise ValueError(f'{path.name}: missing ARD sheet')
        it=wb['ARD'].iter_rows(values_only=True)
        header=next(it,None)
        if header is None: raise ValueError(f'{path.name}: empty ARD')
        cols=[str(c).strip() if c is not None else '' for c in header]
        if len(cols)!=len(set(cols)) or '' in cols: raise ValueError(f'{path.name}: duplicate/blank column headers')
        missing=set(KEYS)-set(cols)
        if missing: raise ValueError(f'{path.name}: missing keys {missing}')
        yield cols,None
        for row in it: yield cols,row
    finally: wb.close()

def main():
    OUTPUT_FOLDER.mkdir(parents=True,exist_ok=True)
    for f in [REFERENCE,*EARLIER]:
        if not (INPUT_FOLDER/f).is_file(): raise FileNotFoundError(INPUT_FOLDER/f)
    summary=[]
    details=OUTPUT_FOLDER/'retention_findings.csv'
    with tempfile.TemporaryDirectory(prefix='quasar_audit_') as tmp:
        db=sqlite3.connect(str(Path(tmp)/'reference.sqlite'))
        try:
            db.execute('PRAGMA journal_mode=DELETE')
            db.execute('CREATE TABLE ref (k TEXT PRIMARY KEY, payload BLOB NOT NULL)')
            refcols=None; refstats=Counter(); pending=[]
            for cols,row in rows_from(INPUT_FOLDER/REFERENCE):
                if row is None:
                    refcols=cols; ix=[cols.index(k) for k in KEYS]; continue
                refstats['rows']+=1
                k=key_for(row,ix)
                if k is None: refstats['incomplete_keys']+=1; continue
                payload={c:canonical(row[i]) for i,c in enumerate(cols) if c not in KEYS and i<len(row) and canonical(row[i]) is not None}
                packed=zlib.compress(json.dumps(payload,ensure_ascii=False,separators=(',',':')).encode(),level=3)
                try:
                    db.execute('INSERT INTO ref VALUES (?,?)',(k,packed))
                except sqlite3.IntegrityError:
                    refstats['duplicate_complete_keys']+=1
                if refstats['rows']%BATCH==0: db.commit()
                if refstats['rows']%10000==0: print(f'W188 indexed {refstats["rows"]:,} rows',flush=True)
            db.commit()
            with details.open('w',newline='',encoding='utf-8-sig') as out:
                w=csv.writer(out)
                w.writerow(['SOURCE_FILE',*KEYS,'COLUMN','STATUS','OLDER_VALUE','W188_VALUE'])
                for fname in EARLIER:
                    stats=Counter(); oldercols=None; old_keys=set()
                    print(f'Comparing {fname}',flush=True)
                    for cols,row in rows_from(INPUT_FOLDER/fname):
                        if row is None:
                            oldercols=cols; ix=[cols.index(k) for k in KEYS]
                            absent=set(cols)-set(refcols)-set(KEYS)
                            stats['columns_absent_in_w188']=len(absent)
                            continue
                        stats['rows']+=1
                        k=key_for(row,ix)
                        if k is None:
                            stats['incomplete_keys']+=1
                            continue
                        if k in old_keys:
                            stats['duplicate_complete_keys']+=1
                            continue
                        old_keys.add(k)
                        found=db.execute('SELECT payload FROM ref WHERE k=?',(k,)).fetchone()
                        vals=json.loads(k)
                        if found is None:
                            stats['keys_missing_in_w188']+=1
                            w.writerow([fname,*vals,'','KEY_MISSING','',''])
                            continue
                        stats['keys_found']+=1
                        current=json.loads(zlib.decompress(found[0]))
                        for i,col in enumerate(cols):
                            if col in KEYS: continue
                            old=canonical(row[i]) if i<len(row) else None
                            if old is None: continue
                            stats['older_populated_cells']+=1
                            if col not in refcols:
                                stats['populated_cells_column_missing']+=1
                                w.writerow([fname,*vals,col,'COLUMN_MISSING',old,''])
                            elif col not in current:
                                stats['populated_cells_now_missing']+=1
                                w.writerow([fname,*vals,col,'VALUE_MISSING',old,''])
                            elif current[col]!=old:
                                stats['populated_cells_changed']+=1
                                w.writerow([fname,*vals,col,'VALUE_CHANGED',old,current[col]])
                            else: stats['populated_cells_equal']+=1
                        if stats['rows']%10000==0: print(f'  compared {stats["rows"]:,} rows',flush=True)
                    summary.append({'SOURCE_FILE':fname,**stats})
                    print('  ',dict(stats),flush=True)
            summary_path=OUTPUT_FOLDER/'retention_summary.csv'
            fields=['SOURCE_FILE','rows','incomplete_keys','duplicate_complete_keys','columns_absent_in_w188','keys_found','keys_missing_in_w188','older_populated_cells','populated_cells_equal','populated_cells_now_missing','populated_cells_column_missing','populated_cells_changed']
            with summary_path.open('w',newline='',encoding='utf-8-sig') as f:
                writer=csv.DictWriter(f,fieldnames=fields);writer.writeheader();writer.writerows(summary)
            print('\nW188:',dict(refstats))
            print('Summary:',summary_path,'\nFindings:',details)
            print('IMPORTANT: Missing/duplicate keys are not reconciled. Review findings before deciding W188 replaces earlier loads.')
        finally: db.close()
if __name__=='__main__': main()
