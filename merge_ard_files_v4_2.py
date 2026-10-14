"""Memory-efficient ARD reconciliation, using SQLite + streaming XLSX I/O.

Dependencies: pip install openpyxl
Clinical matching: USUBJID + AVISIT + AVISITN + AVISIT_ORDER.
Incomplete keys remain separate. Input order determines the first retained value.
Conflicts are preserved as `value1 | value2`, consistent with V3.
"""
from __future__ import annotations
import json
import math
import sqlite3
import tempfile
import gc
import time
import shutil
from collections import Counter
from datetime import date, datetime
from numbers import Number
from pathlib import Path
from openpyxl import load_workbook, Workbook

INPUT_FOLDER = Path(r"C:\Users\JMende95\OneDrive - JNJ\Desktop\ard_data")
OUTPUT_FOLDER = INPUT_FOLDER / "merged_xlsx"
FILES_TO_MERGE = [
    "CNTO1959UCO3001_quasar_ind1_wk32_ard_20260619.xlsx",
    "CNTO1959UCO3001_quasar_ind2_wk32_ard_20260621.xlsx",
    "CNTO1959UCO3001_quasar_maint_wk44_ard_20260616.xlsx",
    "CNTO1959UCO3001_quasar_lte_wk92_ard_20261012.xlsx",
    "CNTO1959UCO3001_quasar_lte_wk140_ard_20261013.xlsx",
    "CNTO1959UCO3001_quasar_lte_wk188_ard_20261013.xlsx",
]
OUTPUT_NAME = "CNTO1959UCO3001_quasar_lte_wk188_ard_merged_20261013.xlsx"
# Split oversized reports across Excel sheets; no clinical records are discarded.
MAX_REPORT_DATA_ROWS = 900000
KEYS = ["USUBJID", "AVISIT", "AVISITN", "AVISIT_ORDER"]
SEP = " | "
STRICT_VISIT_QC = True
BATCH_SIZE = 1000
MAX_EXCEL_COLUMNS = 16384
MAX_EXCEL_ROWS = 1048576


def missing(x):
    return x is None or (isinstance(x, str) and not x.strip()) or (isinstance(x, float) and math.isnan(x))


def display(x):
    if missing(x): return ""
    if isinstance(x, (datetime, date)):
        return x.isoformat(sep=" ") if isinstance(x, datetime) and x.time() != datetime.min.time() else x.strftime("%Y-%m-%d")
    return str(x).strip()


def encode(x):
    if missing(x): return ["missing", None]
    if isinstance(x, Number) and not isinstance(x, bool):
        return ["number", str(float(x))]
    if isinstance(x, (datetime, date)): return ["date", display(x)]
    return ["text", display(x)]


def value_equal(a, b):
    if missing(a) or missing(b): return missing(a) and missing(b)
    if isinstance(a, Number) and not isinstance(a, bool) and isinstance(b, Number) and not isinstance(b, bool):
        return float(a) == float(b)
    return display(a) == display(b)


def pack(x):
    if isinstance(x, datetime): return json.dumps(["datetime", x.isoformat()])
    if isinstance(x, date): return json.dumps(["date", x.isoformat()])
    if isinstance(x, bool): return json.dumps(["bool", x])
    if isinstance(x, Number): return json.dumps(["number", x])
    return json.dumps(["text", str(x)])


def unpack(s):
    typ, value = json.loads(s)
    if typ == "datetime": return datetime.fromisoformat(value)
    if typ == "date": return date.fromisoformat(value)
    return value


def alternatives(x):
    return list(dict.fromkeys(part.strip() for part in display(x).split(SEP) if part.strip()))


def combine(a, b):
    if missing(a): return b
    if missing(b) or value_equal(a, b): return a
    return SEP.join(dict.fromkeys(alternatives(a) + alternatives(b)))


def sheet_rows(ws):
    iterator = ws.iter_rows(values_only=True)
    try: header = next(iterator)
    except StopIteration: raise ValueError(f"Empty sheet: {ws.title}")
    names = [str(v).strip() if v is not None else "" for v in header]
    while names and not names[-1]: names.pop()
    if not names or any(not n for n in names): raise ValueError(f"Blank column header in {ws.title}")
    if len(names) != len(set(names)): raise ValueError(f"Duplicate column headers in {ws.title}")
    for row in iterator:
        values = row[:len(names)]
        if not any(not missing(v) for v in values): continue
        yield names, dict(zip(names, values))


def setup(conn):
    conn.executescript("""
        CREATE TABLE records (rid INTEGER PRIMARY KEY, key_id TEXT UNIQUE NOT NULL,
            k0 TEXT, k1 TEXT, k2 TEXT, k3 TEXT, source TEXT);
        CREATE TABLE cells (rid INTEGER NOT NULL, col TEXT NOT NULL, val TEXT NOT NULL,
            PRIMARY KEY (rid, col));
        CREATE TABLE source_visits (col TEXT NOT NULL, visit TEXT NOT NULL,
            PRIMARY KEY(col, visit));
        CREATE TABLE dictionary (signature TEXT PRIMARY KEY, vals TEXT NOT NULL);
        CREATE TABLE conflicts (stage TEXT, k0 TEXT, k1 TEXT, k2 TEXT, k3 TEXT,
            col TEXT, accumulated_files TEXT, accumulated_value TEXT,
            incoming_file TEXT, incoming_value TEXT, merged_value TEXT);
        CREATE TABLE duplicate_groups (file TEXT, k0 TEXT, k1 TEXT, k2 TEXT, k3 TEXT,
            row_count INTEGER, conflict_columns TEXT);
    """)


def process(conn, file, all_columns, dict_columns, stats):
    print(f"Reading {file.name} ...", flush=True)
    wb = load_workbook(file, read_only=True, data_only=True)
    try:
        if "ARD" not in wb.sheetnames: raise ValueError(f"{file.name}: missing ARD sheet")
        ws = wb["ARD"]
        header = None
        local_groups = {}
        seen_in_file = {}
        count = 0
        incomplete_count = 0
        dup_count = 0
        for names, record in sheet_rows(ws):
            if header is None:
                header = names
                if any(k not in names for k in KEYS): raise ValueError(f"{file.name}: missing required keys")
                if any(n.startswith("__ARD_MERGE_INTERNAL") or n.endswith("__INCOMING__") for n in names):
                    raise ValueError(f"{file.name}: reserved column header")
                for col in names:
                    if col not in all_columns: all_columns.append(col)
            count += 1
            kvals = [record[k] for k in KEYS]
            incomplete = any(missing(v) for v in kvals)
            if incomplete:
                incomplete_count += 1
                key_id = json.dumps(["incomplete", file.name, count])
            else:
                key_id = json.dumps(["complete", [encode(v) for v in kvals]], ensure_ascii=False)
            key_strings = [pack(v) if not missing(v) else None for v in kvals]
            is_local_duplicate = False
            if not incomplete:
                if key_id in seen_in_file:
                    is_local_duplicate = True
                    dup_count += 1
                    if key_id not in local_groups:
                        local_groups[key_id] = [2, set(), kvals]
                    else:
                        local_groups[key_id][0] += 1
                else:
                    seen_in_file[key_id] = count
            found = conn.execute("SELECT rid,source FROM records WHERE key_id=?", (key_id,)).fetchone()
            if found is None:
                cur = conn.execute("INSERT INTO records(key_id,k0,k1,k2,k3,source) VALUES(?,?,?,?,?,?)",
                    (key_id, *key_strings, file.name))
                rid = cur.lastrowid
            else:
                rid = found[0]
            for col, val in record.items():
                if col in KEYS or missing(val): continue
                if "AVISIT" in record and not missing(record["AVISIT"]):
                    conn.execute("INSERT OR IGNORE INTO source_visits VALUES(?,?)", (col, display(record["AVISIT"])))
                previous = conn.execute("SELECT val FROM cells WHERE rid=? AND col=?", (rid, col)).fetchone()
                if previous is None:
                    conn.execute("INSERT INTO cells VALUES(?,?,?)", (rid, col, pack(val)))
                else:
                    old = unpack(previous[0])
                    if not value_equal(old, val):
                        merged = combine(old, val)
                        stage = "WITHIN_FILE_DUPLICATE" if is_local_duplicate else "BETWEEN_FILES"
                        if stage == "WITHIN_FILE_DUPLICATE": local_groups[key_id][1].add(col)
                        conn.execute("INSERT INTO conflicts VALUES(?,?,?,?,?,?,?,?,?,?,?)",
                            (stage, *(display(v) for v in kvals), col, found[1] if found else file.name,
                             display(old), file.name, display(val), display(merged)))
                        conn.execute("UPDATE cells SET val=? WHERE rid=? AND col=?", (pack(merged), rid, col))
            if count % BATCH_SIZE == 0:
                conn.commit()
                print(f"  {count:,} rows", flush=True)
        for _, (n, cols, vals) in local_groups.items():
            conn.execute("INSERT INTO duplicate_groups VALUES(?,?,?,?,?,?,?)",
                (file.name, *(display(v) for v in vals), n, SEP.join(sorted(cols))))
        if "PARAMCD_DICT" in wb.sheetnames:
            for names, record in sheet_rows(wb["PARAMCD_DICT"]):
                for col in names:
                    if col not in dict_columns: dict_columns.append(col)
                sig = json.dumps([encode(record.get(col)) for col in names], ensure_ascii=False)
                conn.execute("INSERT OR IGNORE INTO dictionary VALUES(?,?)", (sig, json.dumps({k:pack(v) for k,v in record.items() if not missing(v)})))
        conn.commit()
        stats.append((file.name, count, len(header or []), incomplete_count, dup_count, len(local_groups)))
        print(f"  Done: {count:,} rows; {dup_count:,} additional within-file duplicates", flush=True)
    finally:
        wb.close()
        del ws
        gc.collect()


def write_sheet(wb, name, headers, iterator):
    ws = wb.create_sheet(name)
    ws.append(headers)
    n = 0
    for row in iterator:
        n += 1
        if n >= MAX_EXCEL_ROWS: raise ValueError(f"Excel row limit exceeded: {name}")
        ws.append(row)
    return n


def write_partitioned_sheet(wb, base_name, headers, iterator, chunk_size=MAX_REPORT_DATA_ROWS):
    """Split reports before the Excel row limit, without dropping records."""
    sheet_no = 1
    ws = wb.create_sheet(base_name)
    ws.append(headers)
    total = 0
    within_sheet = 0
    for row in iterator:
        if within_sheet >= chunk_size:
            sheet_no += 1
            ws = wb.create_sheet(f"{base_name}_{sheet_no}")
            ws.append(headers)
            within_sheet = 0
        ws.append(row)
        within_sheet += 1
        total += 1
    print(f"  {base_name}: {total:,} records across {sheet_no} sheet(s)", flush=True)
    return total


def output(conn, path, all_columns, dict_columns, stats):
    if len(all_columns) > MAX_EXCEL_COLUMNS:
        raise ValueError(f"ARD has {len(all_columns)} columns, exceeding Excel's {MAX_EXCEL_COLUMNS} limit")
    wb = Workbook(write_only=True)
    ordered = KEYS + [col for col in all_columns if col not in KEYS]
    qc = []
    source_visits = {}
    for col, visit in conn.execute("SELECT col,visit FROM source_visits"):
        source_visits.setdefault(col,set()).add(visit)
    output_visits = {}
    def ard_rows():
        cursor = conn.execute("SELECT rid,k0,k1,k2,k3 FROM records ORDER BY rid")
        for rid, *keys in cursor:
            vals = {col:unpack(val) for col,val in conn.execute("SELECT col,val FROM cells WHERE rid=?", (rid,))}
            visit = display(unpack(keys[1])) if keys[1] is not None else ""
            if visit:
                for col in vals: output_visits.setdefault(col,set()).add(visit)
            yield [unpack(k) if k is not None else None for k in keys] + [vals.get(col) for col in ordered[4:]]
    nrows = write_sheet(wb,"ARD",ordered,ard_rows())
    print(f"  ARD: {nrows:,} records", flush=True)
    for col in ordered[4:]:
        src, out = source_visits.get(col,set()), output_visits.get(col,set())
        unexpected = out - src
        qc.append((col, SEP.join(sorted(src)), SEP.join(sorted(out)), SEP.join(sorted(unexpected)), "FAIL" if unexpected else "PASS"))
    fails = [x for x in qc if x[-1] == "FAIL"]
    if STRICT_VISIT_QC and fails:
        raise RuntimeError(f"Visit QC failed for {len(fails)} columns; no output written. Examples: {fails[:5]}")
    def dict_rows():
        for (serialized,) in conn.execute("SELECT vals FROM dictionary"):
            vals = json.loads(serialized)
            yield [unpack(vals[c]) if c in vals else None for c in dict_columns]
    write_partitioned_sheet(wb,"PARAMCD_DICT",dict_columns or ["PARAMCD","PARAM","SOURCE"],dict_rows())
    all_set = set(ordered)
    def comparison_rows():
        for name, rows, cols, incomplete, dups, groups in stats:
            yield [name, rows, rows-dups, cols, incomplete, dups, groups, len(all_set)-cols]
    write_sheet(wb,"FILE_COMPARISON",["FILE","ROWS","ROWS_AFTER_INTERNAL_CONSOLIDATION","COLUMNS","NULL_KEY_ROWS_KEPT_SEPARATE","DUPLICATE_COMPLETE_KEY_ROWS","DUPLICATE_COMPLETE_KEY_GROUPS","MISSING_FROM_GLOBAL_COLUMN_UNION"],comparison_rows())
    conflict_count = conn.execute("SELECT COUNT(*) FROM conflicts").fetchone()[0]
    print(f"  Conflict events: {conflict_count:,} (not necessarily distinct clinical conflicts)", flush=True)
    write_partitioned_sheet(wb,"VALUE_CONFLICTS",KEYS+["CONFLICT_STAGE","COLUMN","ACCUMULATED_FILES","ACCUMULATED_VALUE","INCOMING_FILE","INCOMING_VALUE","MERGED_VALUE"],
        ([*r[1:5],r[0],*r[5:]] for r in conn.execute("SELECT * FROM conflicts")))
    write_partitioned_sheet(wb,"DUPLICATE_KEYS",["FILE",*KEYS,"ROW_COUNT","KEY_STATUS","ACTION","CONFLICT_COLUMNS"],
        ([*r[:6],"COMPLETE","CONSOLIDATED",r[6]] for r in conn.execute("SELECT * FROM duplicate_groups")))
    write_sheet(wb,"COLUMN_VISIT_QC",["COLUMN","SOURCE_POPULATED_VISITS","OUTPUT_POPULATED_VISITS","UNEXPECTED_OUTPUT_VISITS","QC_STATUS"],qc)
    write_sheet(wb,"SOURCE_FILES",["LOAD_ORDER","FILE"],([i,name] for i,name in enumerate(FILES_TO_MERGE,1)))
    path.parent.mkdir(parents=True,exist_ok=True)
    tmp = path.with_name(path.stem + ".partial.xlsx")
    try:
        wb.save(tmp)
        tmp.replace(path)
    finally:
        wb.close()
        gc.collect()
        if tmp.exists():
            try: tmp.unlink()
            except PermissionError: print(f"WARNING: temporary output still locked: {tmp}")
    print(f"Success: {path}\nRows: {nrows:,}; columns: {len(ordered):,}; visit QC failures: {len(fails)}")


def main():
    paths = [INPUT_FOLDER / name for name in FILES_TO_MERGE]
    for path in paths:
        if not path.is_file(): raise FileNotFoundError(path)
    # Windows: avoid WAL sidecar files and ensure connections/cursors are released.
    folder = Path(tempfile.mkdtemp(prefix="ard_merge_"))
    conn = None
    try:
        conn = sqlite3.connect(str(folder / "merge.sqlite"))
        conn.execute("PRAGMA journal_mode=DELETE")
        conn.execute("PRAGMA cache_size=-16000")
        setup(conn)
        columns, dict_columns, stats = [], [], []
        for file in paths:
            process(conn, file, columns, dict_columns, stats)
        output(conn, OUTPUT_FOLDER / OUTPUT_NAME, columns, dict_columns, stats)
        conn.commit()
    finally:
        if conn is not None:
            conn.close()
            del conn
        gc.collect()
        for attempt in range(5):
            try:
                shutil.rmtree(folder)
                break
            except PermissionError:
                if attempt == 4:
                    print(f"WARNING: temporary SQLite files remain locked: {folder}")
                    print("The output workbook, if saved, is unaffected.")
                else:
                    time.sleep(1)

if __name__ == "__main__": main()
