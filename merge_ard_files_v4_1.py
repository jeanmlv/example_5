from __future__ import annotations

import csv
import contextlib
import hashlib
import json
import math
import shutil
import tempfile
from collections import Counter
from datetime import date, datetime
from pathlib import Path

try:
    import duckdb
    import pyarrow as pa
    import pyarrow.parquet as pq
    from openpyxl import Workbook, load_workbook
except ImportError as exc:
    raise SystemExit("Install dependencies: py -m pip install duckdb pyarrow openpyxl") from exc

# -------------------- CONFIGURATION --------------------
INPUT_FOLDER = Path("/mnt/quasar/input")
OUTPUT_FOLDER = Path("/mnt/quasar/output_v4_1")
# Work in the Domino /mnt volume; confirm persistence before closing workspace.
FILES_TO_MERGE = [
    "CNTO1959UCO3001_quasar_ind1_wk32_ard_20260619.xlsx",
    "CNTO1959UCO3001_quasar_ind2_wk32_ard_20260621.xlsx",
    "CNTO1959UCO3001_quasar_maint_wk44_ard_20260616.xlsx",
    "CNTO1959UCO3001_quasar_lte_wk92_ard_20261012.xlsx",
    "CNTO1959UCO3001_quasar_lte_wk140_ard_20261013.xlsx",
    "CNTO1959UCO3001_quasar_lte_wk188_ard_20261013.xlsx",
]
KEYS = ("USUBJID", "AVISIT", "AVISITN", "AVISIT_ORDER")
MIN_FREE_GB = 10  # conservative minimum; large studies may need considerably more
MEMORY_LIMIT = "16GB"  # initial DuckDB limit; container reports ~81 GiB cgroup cap
THREADS = 2
ROWS_PER_PARQUET_BATCH = 250
# CSV is a convenience export only; all conflicts are ALWAYS in Parquet.
MAX_CONFLICT_CSV_ROWS = 900_000
KEEP_STAGING_FILES = False
OBS_PARTITIONS = 128  # hash-based bounded aggregation
# -------------------------------------------------------


def quoted(path: Path) -> str:
    return "'" + str(path).replace("'", "''") + "'"


def normalize(value) -> str:
    """Stable text representation; never silently round clinical numbers."""
    if value is None:
        return ""
    if isinstance(value, str):
        return value.strip()
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, (datetime, date)):
        return value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    if isinstance(value, float):
        if math.isnan(value):
            return ""
        if math.isinf(value):
            return str(value)
        return str(int(value)) if value.is_integer() else repr(value)
    return str(value).strip()


def canonical_key(parts: list[str], filename: str, row_number: int) -> tuple[str, bool]:
    complete = all(bool(p) for p in parts)
    # Incomplete keys NEVER match across rows/files.
    payload = ["COMPLETE", *parts] if complete else ["INCOMPLETE", filename, row_number]
    digest = hashlib.sha256(json.dumps(payload, ensure_ascii=False).encode("utf-8")).hexdigest()
    return digest, complete


def check_disk(folder: Path, required_gb: int = MIN_FREE_GB):
    free = shutil.disk_usage(folder).free / (1024 ** 3)
    print(f"Available disk space at {folder}: {free:.1f} GB", flush=True)
    if free < required_gb:
        raise RuntimeError(
            f"Only {free:.1f} GB free. At least {required_gb} GB is required "
            "before starting. Choose another OUTPUT_FOLDER/drive or free space."
        )


def stream_workbook(path: Path, file_order: int, obs_writer, rows_writer,
                    all_columns: list[str], stats: list[dict]):
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "ARD" not in wb.sheetnames:
            raise ValueError(f"Missing ARD sheet: {path.name}")
        ws = wb["ARD"]
        iterator = ws.iter_rows(values_only=True)
        try:
            headers = [normalize(x) for x in next(iterator)]
        except StopIteration:
            raise ValueError(f"Empty ARD sheet: {path.name}")
        if any(not h for h in headers) or len(set(headers)) != len(headers):
            raise ValueError(f"Blank/duplicate column names in {path.name}")
        if not set(KEYS).issubset(headers):
            raise ValueError(f"Missing clinical keys in {path.name}: {set(KEYS)-set(headers)}")
        for name in headers:
            if name not in all_columns:
                all_columns.append(name)
        key_idx = [headers.index(k) for k in KEYS]
        val_idx = [(i, h) for i, h in enumerate(headers) if h not in KEYS]
        count, populated, incomplete = 0, 0, 0
        for rownum, row in enumerate(iterator, start=2):
            if not row or all(x is None or (isinstance(x, str) and not x.strip()) for x in row):
                continue
            keyparts = [normalize(row[i]) if i < len(row) else "" for i in key_idx]
            rid, complete = canonical_key(keyparts, path.name, rownum)
            rows_writer.writerow([rid, *keyparts, path.name, rownum, file_order, int(complete)])
            count += 1
            incomplete += not complete
            for i, column in val_idx:
                value = normalize(row[i]) if i < len(row) else ""
                if value:
                    obs_writer.writerow([rid, column, value, path.name, file_order, rownum])
                    populated += 1
            if count % 10000 == 0:
                print(f"  {path.name}: {count:,} rows, {populated:,} populated cells", flush=True)
        stats.append({"FILE": path.name, "ROWS": count, "COLUMNS": len(headers),
                      "POPULATED_CELLS": populated, "INCOMPLETE_KEY_ROWS": incomplete})
        print(f"Finished {path.name}: {count:,} rows / {populated:,} populated cells", flush=True)
    finally:
        wb.close()


def export_summary(conn, output_dir: Path, file_stats: list[dict], total_columns: int):
    wb = Workbook(write_only=True)
    sheet = wb.create_sheet("FILE_COMPARISON")
    columns = ["FILE", "ROWS", "COLUMNS", "POPULATED_CELLS", "INCOMPLETE_KEY_ROWS"]
    sheet.append(columns)
    for rec in file_stats:
        sheet.append([rec[c] for c in columns])
    metrics = wb.create_sheet("RECONCILIATION_SUMMARY")
    metrics.append(["METRIC", "VALUE"])
    queries = [
        ("Distinct clinical keys", "SELECT COUNT(*) FROM key_table"),
        ("Variables in union", None),
        ("Consolidated populated cells", "SELECT COUNT(*) FROM cells"),
        ("Cells with conflicts", "SELECT COUNT(*) FROM cells WHERE n_values > 1"),
        ("Distinct value alternatives", "SELECT COUNT(*) FROM distinct_values"),
        ("Keys present in multiple input rows", "SELECT COUNT(*) FROM key_table WHERE input_rows > 1"),
        ("Keys found in multiple source files", "SELECT COUNT(*) FROM key_table WHERE n_files > 1"),
    ]
    for label, query in queries:
        val = total_columns if query is None else conn.execute(query).fetchone()[0]
        metrics.append([label, val])
    notes = wb.create_sheet("IMPORTANT_NOTES")
    for text in [
        "Values with >1 distinct alternative are represented as JSON arrays in the merged ARD.",
        "Conflict cells are NOT numeric measurements; consult the lineage Parquet before analysis.",
        "All original nonempty cell alternatives and their source files are in DATA_LINEAGE.parquet.",
        "Only complete 4-component clinical keys are matched; incomplete keys stay separate.",
        "This audit cannot establish whether source values were observed or imputed.",
        "No automatic choice of latest source value is made when values disagree.",
        "Excel exports are summary-only; full ARD is in Parquet to avoid worksheet row limits.",
    ]:
        notes.append([text])
    wb.save(output_dir / "QUASAR_RECONCILIATION_SUMMARY.xlsx")


def write_wide_parquet(conn, path: Path, all_columns: list[str]):
    """Stream sorted sparse cells into wide Parquet without pivoting in RAM."""
    schema = pa.schema([(col, pa.string()) for col in all_columns])
    writer = pq.ParquetWriter(path, schema, compression="zstd")
    sql = """
        SELECT k.key_id, k.USUBJID, k.AVISIT, k.AVISITN, k.AVISIT_ORDER,
               c.col_name, c.merged_value, c.n_values
        FROM key_table k LEFT JOIN cells c ON k.key_id = c.key_id
        ORDER BY k.key_id, c.col_name
    """
    cursor = conn.execute(sql)
    output_batch: list[dict] = []
    current_id = None
    current_row = None
    written = 0
    try:
        while True:
            fetched = cursor.fetchmany(20000)
            if not fetched:
                break
            for rid, subject, visit, visitn, order, col, value, n_values in fetched:
                if rid != current_id:
                    if current_row is not None:
                        output_batch.append(current_row)
                    current_id = rid
                    current_row = {"USUBJID": subject, "AVISIT": visit,
                                   "AVISITN": visitn, "AVISIT_ORDER": order}
                if col is not None:
                    # JSON array is unambiguous even when original values contain '|'.
                    current_row[col] = value
                if len(output_batch) >= ROWS_PER_PARQUET_BATCH:
                    table = pa.Table.from_pylist(output_batch, schema=schema)
                    writer.write_table(table)
                    written += len(output_batch)
                    output_batch.clear()
                    if written % 5000 < ROWS_PER_PARQUET_BATCH:
                        print(f"  Wrote {written:,} merged ARD rows", flush=True)
        if current_row is not None:
            output_batch.append(current_row)
        if output_batch:
            writer.write_table(pa.Table.from_pylist(output_batch, schema=schema))
            written += len(output_batch)
    finally:
        writer.close()
    print(f"Wide ARD written: {written:,} rows, {len(all_columns):,} columns", flush=True)


def partition_observations(obs_csv: Path, folder: Path, n: int):
    """Partition by key_id so every value for a clinical key stays together.

    Only one output handle is open at a time to avoid OS file descriptor limits.
    """
    folder.mkdir(parents=True, exist_ok=True)
    paths = [folder / f"obs_{i:03d}.csv" for i in range(n)]
    # Buffer one modest chunk in RAM; distribute by first hash digits.
    buffers = [[] for _ in range(n)]
    buffered = 0
    total = 0

    def flush():
        nonlocal buffered
        for i, records in enumerate(buffers):
            if not records:
                continue
            with paths[i].open("a", newline="", encoding="utf-8") as out:
                csv.writer(out).writerows(records)
            records.clear()
        buffered = 0

    with obs_csv.open(newline="", encoding="utf-8") as inp:
        reader = csv.reader(inp)
        header = next(reader)
        for i, path in enumerate(paths):
            with path.open("w", newline="", encoding="utf-8") as out:
                csv.writer(out).writerow(header)
        for rec in reader:
            bucket = int(rec[0][:8], 16) % n
            buffers[bucket].append(rec)
            buffered += 1
            total += 1
            if buffered >= 10000:
                flush()
            if total % 5000000 == 0:
                print(f"  Partitioned {total:,} observations", flush=True)
    flush()
    print(f"Partitioning finished: {total:,} observations in {n} buckets", flush=True)
    return paths


def aggregate_buckets(conn, paths: list[Path], stage: Path):
    distinct_dir = stage / "distinct_parts"
    cells_dir = stage / "cell_parts"
    distinct_dir.mkdir()
    cells_dir.mkdir()
    for i, path in enumerate(paths):
        distinct_path = distinct_dir / f"part_{i:03d}.parquet"
        cells_path = cells_dir / f"part_{i:03d}.parquet"
        # Each partition is a disjoint set of key_ids. No global GROUP BY needed.
        conn.execute(f"""
            CREATE TEMP TABLE bucket_distinct AS
            SELECT key_id, col_name, val,
                   MIN(TRY_CAST(file_order AS INTEGER)) AS first_order,
                   STRING_AGG(DISTINCT source_file, ' | ' ORDER BY source_file) AS source_files,
                   COUNT(*) AS source_occurrences
            FROM read_csv({quoted(path)}, header=true, all_varchar=true)
            GROUP BY key_id, col_name, val
        """)
        conn.execute("""
            CREATE TEMP TABLE bucket_cells AS
            SELECT key_id, col_name, COUNT(*) AS n_values,
                   CASE WHEN COUNT(*)=1 THEN MIN(val)
                        ELSE to_json(list(val ORDER BY first_order, val)) END AS merged_value
            FROM bucket_distinct
            GROUP BY key_id, col_name
        """)
        conn.execute(f"COPY bucket_distinct TO {quoted(distinct_path)} (FORMAT PARQUET, COMPRESSION ZSTD)")
        conn.execute(f"COPY bucket_cells TO {quoted(cells_path)} (FORMAT PARQUET, COMPRESSION ZSTD)")
        conn.execute("DROP TABLE bucket_cells")
        conn.execute("DROP TABLE bucket_distinct")
        path.unlink()
        if (i + 1) % 8 == 0 or i + 1 == len(paths):
            print(f"  Aggregated {i+1}/{len(paths)} buckets", flush=True)
            check_disk(OUTPUT_FOLDER, required_gb=2)
    conn.execute(f"CREATE VIEW distinct_values AS SELECT * FROM read_parquet({quoted(distinct_dir / '*.parquet')})")
    conn.execute(f"CREATE VIEW cells AS SELECT * FROM read_parquet({quoted(cells_dir / '*.parquet')})")


def main():
    for filename in FILES_TO_MERGE:
        if not (INPUT_FOLDER / filename).is_file():
            raise FileNotFoundError(INPUT_FOLDER / filename)
    OUTPUT_FOLDER.mkdir(parents=True, exist_ok=True)
    check_disk(OUTPUT_FOLDER)
    stage = OUTPUT_FOLDER / "_working_v4_1"
    if stage.exists():
        raise RuntimeError(f"Staging directory already exists: {stage}. "
                           "Review/remove it manually before a new run.")
    stage.mkdir()
    obs_csv = stage / "observations.csv"
    rows_csv = stage / "row_keys.csv"
    file_stats: list[dict] = []
    all_columns = list(KEYS)
    try:
        with obs_csv.open("w", newline="", encoding="utf-8") as fo, \
             rows_csv.open("w", newline="", encoding="utf-8") as fr:
            ow, rw = csv.writer(fo), csv.writer(fr)
            ow.writerow(["key_id", "col_name", "val", "source_file", "file_order", "row_number"])
            rw.writerow(["key_id", *KEYS, "source_file", "row_number", "file_order", "complete_key"])
            for order, filename in enumerate(FILES_TO_MERGE, start=1):
                stream_workbook(INPUT_FOLDER / filename, order, ow, rw, all_columns, file_stats)
                check_disk(OUTPUT_FOLDER, required_gb=2)
        print("Aggregating with DuckDB...", flush=True)
        db = stage / "reconciliation.duckdb"
        conn = duckdb.connect(str(db))
        try:
            conn.execute(f"SET memory_limit='{MEMORY_LIMIT}'")
            conn.execute(f"SET threads={THREADS}")
            conn.execute(f"SET temp_directory={quoted(stage / 'duckdb_tmp')}")
            conn.execute("SET preserve_insertion_order=false")
            conn.execute(f"""
                CREATE TABLE key_table AS
                SELECT key_id,
                       MIN(USUBJID) AS USUBJID, MIN(AVISIT) AS AVISIT,
                       MIN(AVISITN) AS AVISITN, MIN(AVISIT_ORDER) AS AVISIT_ORDER,
                       COUNT(*) AS input_rows, COUNT(DISTINCT source_file) AS n_files
                FROM read_csv({quoted(rows_csv)}, header=true, all_varchar=true)
                GROUP BY key_id
            """)
            # Bounded-memory reconciliation: no global GROUP BY on all observations.
            print("Partitioning observations by clinical key...", flush=True)
            obs_parts = partition_observations(obs_csv, stage / "obs_parts", OBS_PARTITIONS)
            print("Aggregating partitions with DuckDB...", flush=True)
            aggregate_buckets(conn, obs_parts, stage)
            lineage = OUTPUT_FOLDER / "QUASAR_DATA_LINEAGE.parquet"
            conn.execute(f"COPY (SELECT * FROM distinct_values) TO {quoted(lineage)} (FORMAT PARQUET, COMPRESSION ZSTD)")
            conflict_parquet = OUTPUT_FOLDER / "QUASAR_VALUE_CONFLICTS.parquet"
            conn.execute(f"""
                COPY (
                    SELECT k.USUBJID, k.AVISIT, k.AVISITN, k.AVISIT_ORDER,
                           d.key_id, d.col_name AS COLUMN, d.val AS ALTERNATIVE_VALUE,
                           d.source_files AS SOURCE_FILES, d.source_occurrences,
                           c.n_values AS DISTINCT_VALUES
                    FROM distinct_values d
                    JOIN cells c ON d.key_id=c.key_id AND d.col_name=c.col_name
                    JOIN key_table k ON k.key_id=d.key_id
                    WHERE c.n_values > 1
                ) TO {quoted(conflict_parquet)} (FORMAT PARQUET, COMPRESSION ZSTD)
            """)
            n_conflict_rows = conn.execute(f"""
                SELECT COALESCE(SUM(n_values),0) FROM cells WHERE n_values>1
            """).fetchone()[0]
            if n_conflict_rows <= MAX_CONFLICT_CSV_ROWS:
                csv_path = OUTPUT_FOLDER / "QUASAR_VALUE_CONFLICTS.csv"
                conn.execute(f"""
                    COPY (
                        SELECT k.USUBJID, k.AVISIT, k.AVISITN, k.AVISIT_ORDER,
                               d.col_name AS COLUMN, d.val AS ALTERNATIVE_VALUE,
                               d.source_files AS SOURCE_FILES
                        FROM distinct_values d
                        JOIN cells c ON d.key_id=c.key_id AND d.col_name=c.col_name
                        JOIN key_table k ON k.key_id=d.key_id
                        WHERE c.n_values > 1
                    ) TO {quoted(csv_path)} (HEADER, DELIMITER ',')
                """)
            else:
                print(f"Conflict CSV omitted ({n_conflict_rows:,} alternatives); "
                      "complete conflict report is in Parquet.", flush=True)
            # Row-level duplicate report is compact; no arbitrary truncation.
            conn.execute(f"""
                COPY (
                    SELECT * FROM key_table WHERE input_rows>1
                ) TO {quoted(OUTPUT_FOLDER / 'QUASAR_DUPLICATE_KEYS.parquet')}
                (FORMAT PARQUET, COMPRESSION ZSTD)
            """)
            export_summary(conn, OUTPUT_FOLDER, file_stats, len(all_columns))
            write_wide_parquet(conn, OUTPUT_FOLDER / "QUASAR_ARD_MERGED.parquet", all_columns)
        finally:
            conn.close()
        print("SUCCESS: review QUASAR_RECONCILIATION_SUMMARY.xlsx and conflicts before analysis.")
    finally:
        if not KEEP_STAGING_FILES:
            try:
                shutil.rmtree(stage)
            except OSError as exc:
                print(f"WARNING: Could not remove staging directory {stage}: {exc}")


if __name__ == "__main__":
    main()
